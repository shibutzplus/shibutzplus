"""
Converter for "Yitzhak Sadeh A Bat Yam" School Schedules
School ID: xu77qkr9250ki6g77c9k3iyd

Converts:
  1) מערכת שעות לכיתה - טבלה.pdf (+ שינויים) -> כיתות.docx
  2) מערכת שעות למורה - טבלה.pdf (+ שינויים) -> מורים.docx

Formatted in the exact Korczak style for seamless import into ShibutzPlus.
Automatically merges incremental changes from 'שינויים' directory.
"""

import os
import sys
import re
from datetime import datetime
import pymupdf
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

if sys.stdout:
    sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
CLASS_PDF = os.path.join(BASE_DIR, "מערכת שעות לכיתה - טבלה.pdf")
TEACHER_PDF = os.path.join(BASE_DIR, "מערכת שעות למורה - טבלה.pdf")
CLASS_DOCX = os.path.join(BASE_DIR, "כיתות.docx")
TEACHER_DOCX = os.path.join(BASE_DIR, "מורים.docx")

# Changes directory and batch files (in chronological order)
SHINUYIM_DIR = os.path.join(BASE_DIR, "שינויים")
BATCHES = [
    {
        "name": "Batch 1 (28/09/2026)",
        "teacher_pdf": os.path.join(SHINUYIM_DIR, "מערכת שעות למורה - טבלה (1).pdf"),
        "class_pdf": os.path.join(SHINUYIM_DIR, "מערכת שעות לכיתה - טבלה (1).pdf"),
    },
    {
        "name": "Batch 2 (04/10/2026)",
        "teacher_pdf": os.path.join(SHINUYIM_DIR, "מערכת שעות למורה - טבלה.pdf"),
        "class_pdf": os.path.join(SHINUYIM_DIR, "מערכת שעות לכיתה - טבלה.pdf"),
    },
]

# School metadata
SCHOOL_ID = "xu77qkr9250ki6g77c9k3iyd"
SCHOOL_NAME = "יצחק שדה א"
META_TIMESTAMP = "04/10/2026 18:00:00"

HOUR_LABELS = {
    0: "שעה 0",
    1: "שעה 1 (08:00-08:50)",
    2: "שעה 2 (08:50-09:35)",
    3: "שעה 3 (10:15-11:00)",
    4: "שעה 4 (11:00-11:45)",
    5: "שעה 5 (12:00-12:45)",
    6: "שעה 6 (12:45-13:30)",
    7: "שעה 7 (13:30-14:20)",
    8: "שעה 8",
    9: "שעה 9",
    10: "שעה 10",
}

DAYS_ORDER = [
    (1, "יום ראשון"),
    (2, "יום שני"),
    (3, "יום שלישי"),
    (4, "יום רביעי"),
    (5, "יום חמישי"),
    (6, "יום שישי")
]

# Map PDF column index to Day number (1=Sunday, 6=Friday)
# Column 5: ראשון, 4: שני, 3: שלישי, 2: רביעי, 1: חמישי, 0: שישי
COL_TO_DAY = {
    5: 1,  # ראשון
    4: 2,  # שני
    3: 3,  # שלישי
    2: 4,  # רביעי
    1: 5,  # חמישי
    0: 6   # שישי
}

# Homeroom teachers mapping
CLASS_HOMEROOM = {
    'א1': 'עינב מזרחי',
    'א2': 'דורין נגדי',
    'א3': 'דניאל קדושי בדוסה',
    'א4': 'סימה קהלאני',
    'ב1': 'נוי חזן גרמן',
    'ב2': 'לירון פינטו עובד',
    'ב3': 'נוי חזן גרמן',
    'ב4': 'מאי נאמן',
    'ג1': 'שנהב רחמני',
    'ג2': 'הילה ציון',
    'ג3': 'ענבל חי',
    'ג4': 'לינוי בן משה',
    'ד1': 'ילנה דרמנסקי',
    'ד2': 'שיר טויג',
    'ד3': 'קרן שמעוני',
    'ד4': 'עדי פנחס',
    'ד5': 'אופיר יצחק',
    'ה1': 'יעל קייזר',
    'ה2': 'לימור אברגל',
    'ה3': 'הילה כהן ישר',
    'ה4': 'מריל שמואל',
    'ו1': 'חגית אלבז',
    'ו2': 'נינה איילון',
    'ו3': 'דניאלה ששון דניאל',
    'ו4': 'קארן ונגרוביץ טרבלוס'
}

# Map raw teacher names from PDF to canonical First Last
TEACHER_NAME_MAP = {
    'אבזוב יוליה': 'יוליה אבזוב',
    'אברגל לימור': 'לימור אברגל',
    'איילון נינה': 'נינה איילון',
    'אלבז חגית': 'חגית אלבז',
    'אלמולי טלי': 'טלי אלמולי',
    'אנדורן סער חגית': 'חגית אנדורן סער',
    'בן משה לינוי': 'לינוי בן משה',
    'לינוי בן משה': 'לינוי בן משה',
    'משה לינוי-בן': 'לינוי בן משה',
    'משה לינוי': 'לינוי בן משה',
    'ברגיג ליטל': 'ליטל ברגיג',
    'ברזילי מעיין': 'מעיין ברזילי',
    'גרודן יונה': 'יונה גרודן',
    'דיין בוחניק עדן': 'עדן דיין בוחניק',
    'דרמנסקי ילנה': 'ילנה דרמנסקי',
    'ונגרוביץ טרבלוס קארן': 'קארן ונגרוביץ טרבלוס',
    'ונגרוביץ טרבלוס קאר': 'קארן ונגרוביץ טרבלוס',
    'חודוס אלכסנדרה': 'אלכסנדרה חודוס',
    'חזן גרמן נוי': 'נוי חזן גרמן',
    'חי ענבל': 'ענבל חי',
    'חלימוב מרטין': 'מרטין חלימוב',
    'חתו נפין': 'נפין חתו',
    'טויג שיר': 'שיר טויג',
    'טויזר סיגלית': 'סיגלית טויזר',
    'יוסף מעין': 'מעין יוסף',
    'יצחק אופיר': 'אופיר יצחק',
    'כהן ישר הילה': 'הילה כהן ישר',
    'כהן ספיר': 'ספיר כהן',
    'ליאון ויקטוריה': 'ויקטוריה ליאון',
    'לסרי מירב': 'מירב לסרי',
    'לרנר יהודית': 'יהודית לרנר',
    'מדריך גאלינג': 'גאלינג מדריך',
    'גאלינג מדריך': 'גאלינג מדריך',
    'מדריך הפרעות נפשיות': 'מדריך הפרעות נפשיות',
    'מדריך טכנולגו': 'טכנולוגיה מדריך',
    'טכנולוגיה מדריך': 'טכנולוגיה מדריך',
    'מדריך שחמט': 'שחמט מדריך',
    'שחמט מדריך': 'שחמט מדריך',
    'מדריך ASD': 'ASD מדריך',
    'ASD מדריך': 'ASD מדריך',
    'אוראל מהצרי לב': 'אוראל מהצרי לב',
    'מהצרי לב אוראל': 'אוראל מהצרי לב',
    'מורה לקולנוע': 'מורה לקולנוע',
    'מורה רובוטק': 'מורה רובוטק',
    'מזרחי עינב': 'עינב מזרחי',
    'מנצור שגיא צדוק': 'שגיא צדוק מנצור',
    'נאמן מאי': 'מאי נאמן',
    'נגדי דורין': 'דורין נגדי',
    'זילברברג אורית שמחה': 'אורית שמחה זילברברג',
    'זילברברג אורית שמ': 'אורית שמחה זילברברג',
    'סטרו רויטל': 'רויטל סטרו',
    'פאר בובליל עינת': 'עינת פאר בובליל',
    'פינטו עובד לירון': 'לירון פינטו עובד',
    'פנחס עדי': 'עדי פנחס',
    'פקרש יוליה': 'יוליה פקרש',
    'פרי איילת': 'איילת פרי',
    'ציון הילה': 'הילה ציון',
    'קדושי בדוסה דניאל': 'דניאל קדושי בדוסה',
    'קהלאני סימה': 'סימה קהלאני',
    'קייזר יעל': 'יעל קייזר',
    'קמחי מיכל': 'מיכל קמחי',
    'רחמני שנהב': 'שנהב רחמני',
    'שמואל מריל': 'מריל שמואל',
    'שמעוני קרן': 'קרן שמעוני',
    'ששון דניאל דניאלה': 'דניאלה ששון דניאל',
    'מורה תאטרון': 'מורה תאטרון',
}

MULTI_CLASS_MAP = {
    "1 ' ,ד3 ' ד": ['ד1', 'ד3'],
    "3 ' ,ג1 ' ג": ['ג1', 'ג3'],
    "3 ' ,ד1 ' ד": ['ד1', 'ד3'],
    "4 ' ,א3 ' ,א2 ' א": ['א2', 'א3', 'א4'],
    "4 ' ,ב3 ' ,ב2 ' ב": ['ב2', 'ב3', 'ב4'],
    "4 ' ,ג3 ' ,ג2 ' ג": ['ג2', 'ג3', 'ג4'],
    "4 ' ,ד3 ' ,ד2 ' ד": ['ד2', 'ד3', 'ד4'],
    "4 ' ,ד5 ' ד": ['ד4', 'ד5'],
    "5 ' ,ד4 ' ד": ['ד4', 'ד5'],
    "4 ' ,ה3 ' ,ה2 ' ה": ['ה2', 'ה3', 'ה4'],
    "4 ' ,ו3 ' ,ו2 ' ו": ['ו2', 'ו3', 'ו4'],
    "1 ' א": ['א1'], "2 ' א": ['א2'], "3 ' א": ['א3'], "4 ' א": ['א4'],
    "1 ' ב": ['ב1'], "2 ' ב": ['ב2'], "3 ' ב": ['ב3'], "4 ' ב": ['ב4'],
    "1 ' ג": ['ג1'], "2 ' ג": ['ג2'], "3 ' ג": ['ג3'], "4 ' ג": ['ג4'],
    "1 ' ד": ['ד1'], "2 ' ד": ['ד2'], "3 ' ד": ['ד3'], "4 ' ד": ['ד4'], "5 ' ד": ['ד5'],
    "1 ' ה": ['ה1'], "2 ' ה": ['ה2'], "3 ' ה": ['ה3'], "4 ' ה": ['ה4'],
    "1 ' ו": ['ו1'], "2 ' ו": ['ו2'], "3 ' ו": ['ו3'], "4 ' ו": ['ו4'],
}

SUBJECT_CLEANUP = {
    'מליאת מורים': 'מליאה',
    'מליאה': 'מליאה',
    'מקהלת בית הספר': 'מקהלה',
    'מקהלה': 'מקהלה',
    'שעות ייעוץ': 'ייעוץ',
    'ייעוץ': 'ייעוץ',
    'שעות שילוב': 'שילוב',
    'שילוב': 'שילוב',
    'שעות ניהול': 'ניהול',
    'ניהול': 'ניהול',
    'צוות חנ"מ,שעות ייעו': 'צוות חנמ / ייעוץ',
    'צוות חנ"מ / שעות ייעוץ': 'צוות חנמ / ייעוץ',
    'צוות חנמ / שעות ייעוץ': 'צוות חנמ / ייעוץ',
    'צוות חנמ / ייעוץ': 'צוות חנמ / ייעוץ',
    'פיתוח הון א- תפקיד': 'פיתוח הון אנושי',
    'פיתוח הון אנושי - תפקיד': 'פיתוח הון אנושי',
    'פיתוח הון אנושי': 'פיתוח הון אנושי',
    "1'צוות חנ\"מ כיתה א": 'צוות חנמ כיתה א',
    "1'צוות חנ\"מ כיתה ב": 'צוות חנמ כיתה ב',
    "1'צוות חנ\"מ כיתה ג": 'צוות חנמ כיתה ג',
    "1'צוות חנ\"מ כיתה ד": 'צוות חנמ כיתה ד',
    "1'צוות חנ\"מ כיתה ה": 'צוות חנמ כיתה ה',
    "1'צוות חנ\"מ כיתה ו": 'צוות חנמ כיתה ו',
    "5'צוות חנ\"מ כיתה ד": 'צוות חנמ כיתה ד5',
    'צוות חנ"מ כיתה א': 'צוות חנמ כיתה א',
    'צוות חנ"מ כיתה ב': 'צוות חנמ כיתה ב',
    'צוות חנ"מ כיתה ג': 'צוות חנמ כיתה ג',
    'צוות חנ"מ כיתה ד': 'צוות חנמ כיתה ד',
    'צוות חנ"מ כיתה ד5': 'צוות חנמ כיתה ד5',
    'צוות חנ"מ כיתה ה': 'צוות חנמ כיתה ה',
    'צוות חנ"מ כיתה ו': 'צוות חנמ כיתה ו',
    'קבוצות שפה שכבת': 'קבוצות שפה',
    'קבוצות שפה שכבת ו': 'קבוצות שפה שכבת ו',
    'הדרכת הפרעות נפש': 'הדרכת הפרעות נפשיות',
    'הדרכת תסמונות נדי': 'הדרכת תסמונות נדירות',
    'ב-צוות שכבות א': 'צוות שכבות א-ב',
    'ד-צוות שכבות ג': 'צוות שכבות ג-ד',
    'ו-צוות שכבות ה': 'צוות שכבות ה-ו',
    'ASD הדרכת': 'הדרכת ASD',
    'טכנולגו': 'טכנולוגיה',
}

WORKGROUP_SUBJECTS = {
    'פרטני',
    'ניהול כיתה',
    'ניהול',
    'שעות ניהול',
    'קבוצות שפה',
    'קבוצות שפה שכבת ו',
    'קבוצות מתמטיקה',
    'שהייה',
    'שילוב',
    'שעות שילוב',
    'הוראה מותאמת',
    'מיומנויות חברתיות',
    'ייעוץ',
    'שעות ייעוץ',
    'יועצת ומנהל',
    'מליאה',
    'מליאת מורים',
    'מקהלה',
    'מקהלת בית הספר',
    'קבוצת העצמה',
    'ניתוח התנהגות',
    'פיתוח הון אנושי',
    'פיתוח הון אנושי - תפקיד',
    'צוות ניהול מצומצם',
    'צוות שכבת א',
    'צוות שכבת ב',
    'צוות שכבת ג',
    'צוות שכבת ד',
    'צוות שכבת ה',
    'צוות שכבת ו',
    'צוות שכבות א-ב',
    'צוות שכבות ג-ד',
    'צוות שכבות ה-ו',
    'צוות אנגלית',
    'צוות חנמ כיתה א',
    'צוות חנמ כיתה ב',
    'צוות חנמ כיתה ג',
    'צוות חנמ כיתה ד',
    'צוות חנמ כיתה ה',
    'צוות חנמ כיתה ו',
    'צוות חנמ כיתה ד5',
    'צוות חנמ / ייעוץ',
    'הדרכת ASD',
    'הדרכת הפרעות נפשיות',
    'הדרכת תסמונות נדירות',
    'הדרכה יהודית לרנר',
    'משעולים',
    'נבחרת',
}


def is_workgroup(name):
    if not name:
        return False
    name_clean = clean_subject(name)
    if name_clean in WORKGROUP_SUBJECTS:
        return True
    return any(k in name_clean for k in ['פרטני', 'ניהול כיתה', 'ניהול', 'שהייה', 'קבוצות שפה', 'קבוצות מתמטיקה', 'שילוב', 'צוות', 'הדרכ', 'ייעוץ', 'מליא', 'מקהל', 'משעול', 'נבחרת'])


def clean_subject(raw_subj):
    if not raw_subj:
        return ""
    s = raw_subj.strip()
    return SUBJECT_CLEANUP.get(s, s)


def clean_teacher(name):
    if not name:
        return ""
    clean = name.strip()
    return TEACHER_NAME_MAP.get(clean, clean)


def parse_teacher_title(title_line):
    clean = title_line.strip()
    if 'משה לינוי-מערכת שעות למורה בן' in clean:
        return 'לינוי בן משה'
    if 'לב אוראל-מערכת שעות למורה מהצרי' in clean:
        return 'אוראל מהצרי לב'
    if 'ASD מערכת שעות למורה מדריך' in clean:
        return 'ASD מדריך'
    clean = re.sub(r'מערכת שעות\s+(?:ל?מורה|מורה:?)\s*', '', clean).strip()
    return clean_teacher(clean)


def parse_class_title(title_line):
    m = re.search(r'([1-5])\s*[\'"]?\s*מערכת שעות כיתה\s*([א-ו])', title_line)
    if m:
        return f"{m.group(2)}{m.group(1)}"
    m2 = re.search(r'([1-5])\s*[\'"]?\s*([א-ו])', title_line)
    if m2:
        return f"{m2.group(2)}{m2.group(1)}"
    return ""


def format_class_reference(code):
    """Formats a single class code with its homeroom teacher: e.g. 'ד2 שיר טויג'"""
    hr = CLASS_HOMEROOM.get(code, "")
    return f"{code} {hr}".strip() if hr else code


def format_classes_string(raw_classes_str):
    """
    Parses class references from teacher schedule cells.
    Handles single and multi-classes.
    """
    s = raw_classes_str.strip()
    classes = MULTI_CLASS_MAP.get(s, None)
    if classes:
        return ",".join(format_class_reference(c) for c in classes)
    m = re.search(r'([1-5])\s*[\'"]?\s*([א-ו])', s)
    if m:
        code = f"{m.group(2)}{m.group(1)}"
        return format_class_reference(code)
    m2 = re.search(r'([א-ו])\s*[\'"]?\s*([1-5])', s)
    if m2:
        code = f"{m2.group(1)}{m2.group(2)}"
        return format_class_reference(code)
    return s


def set_cell_rtl(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(r'<w:tcMar %s><w:top w:w="80" w:type="dxa"/><w:bottom w:w="80" w:type="dxa"/><w:left w:w="120" w:type="dxa"/><w:right w:w="120" w:type="dxa"/></w:tcMar>' % nsdecls('w'))
    tcPr.append(tcMar)
    for p in cell.paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        pPr = p._p.get_or_add_pPr()
        pPr.append(parse_xml(r'<w:bidi %s/>' % nsdecls('w')))
        for r in p.runs:
            rPr = r._r.get_or_add_rPr()
            rPr.append(parse_xml(r'<w:rtl %s/>' % nsdecls('w')))


def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)


def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="none" w:sz="4"/>'
        f'  <w:bottom w:val="none" w:sz="4"/>'
        f'  <w:left w:val="none" w:sz="4"/>'
        f'  <w:right w:val="none" w:sz="4"/>'
        f'  <w:insideH w:val="none" w:sz="4"/>'
        f'  <w:insideV w:val="none" w:sz="4"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)


def extract_teachers_from_pdf(pdf_path):
    """
    Parses a teacher PDF file.
    Returns:
      schedules: dict of canonical_teacher -> {'teacher', 'day_lessons', 'class_lessons'}
      teacher_order: list of teacher names in order of pages
    """
    doc_pdf = pymupdf.open(pdf_path)
    schedules = {}
    teacher_order = []

    for pno in range(len(doc_pdf)):
        page = doc_pdf[pno]
        lines = [l.strip() for l in page.get_text().split('\n') if l.strip()]
        title_line = [l for l in lines if 'מערכת שעות' in l][0]
        canonical_teacher = parse_teacher_title(title_line)
        teacher_order.append(canonical_teacher)

        tabs = page.find_tables()
        if not tabs.tables:
            continue
        t = tabs.tables[0]

        t_day_lessons = {d_num: [] for d_num in range(1, 7)}
        class_lessons = []

        for r_idx in range(2, len(t.rows)):
            row = t.rows[r_idx]
            h_cell = row.cells[6] if len(row.cells) > 6 else None
            h_text = page.get_text('text', clip=h_cell).strip() if h_cell else ''
            h_num = int(h_text) if h_text.isdigit() else 0

            for col_idx, day_num in COL_TO_DAY.items():
                cell = row.cells[col_idx]
                if not cell:
                    continue
                txt = page.get_text('text', clip=cell).strip()
                if not txt:
                    continue
                c_lines = [l.strip() for l in txt.split('\n') if l.strip()]
                l_type = c_lines[0]

                if len(c_lines) == 1:
                    if l_type == 'פרטני':
                        lesson_str = "פרטני, פרטני"
                    elif l_type == 'תפקיד':
                        lesson_str = "תפקיד, תפקיד"
                    else:
                        lesson_str = f"{l_type}, {l_type}"
                    t_day_lessons[day_num].append((h_num, lesson_str))

                elif len(c_lines) == 2:
                    subj = clean_subject(c_lines[1])
                    if subj == 'פרטני':
                        lesson_str = "פרטני, פרטני"
                    elif subj in ['שעות ניהול', 'ניהול']:
                        lesson_str = "ניהול, תפקיד"
                    elif l_type in ['שהייה', 'תפקיד', 'פרטני']:
                        lesson_str = f"{subj}, {l_type}"
                    else:
                        lesson_str = f"{subj}, שהייה"
                    t_day_lessons[day_num].append((h_num, lesson_str))

                elif len(c_lines) >= 3:
                    raw_cls = c_lines[1]
                    subj = clean_subject(c_lines[2])

                    if subj == 'פרטני':
                        lesson_str = "פרטני, פרטני"
                    elif subj == 'ניהול כיתה':
                        lesson_str = "ניהול כיתה, תפקיד"
                    elif subj in ['קבוצות שפה', 'קבוצות מתמטיקה']:
                        lesson_str = f"{subj}, קבוצה"
                    elif is_workgroup(subj):
                        lesson_str = f"{subj}, {l_type if l_type in ['שהייה', 'תפקיד', 'פרטני'] else 'שהייה'}"
                    else:
                        cls_str = format_classes_string(raw_cls)
                        lesson_str = f"{subj}, {cls_str}, {l_type}"

                        classes = MULTI_CLASS_MAP.get(raw_cls, [])
                        for c_code in classes:
                            class_lessons.append((c_code, day_num, h_num, {
                                'teacher': canonical_teacher,
                                'subject': subj,
                                'type': l_type
                            }))

                    t_day_lessons[day_num].append((h_num, lesson_str))

        schedules[canonical_teacher] = {
            'teacher': canonical_teacher,
            'day_lessons': t_day_lessons,
            'class_lessons': class_lessons
        }

    return schedules, teacher_order


def create_teachers_docx(teacher_schedules):
    print(f"Creating Teachers DOCX: {TEACHER_DOCX}")
    doc_out = Document()

    for section in doc_out.sections:
        section.top_margin = Inches(0.24)
        section.bottom_margin = Inches(0.48)
        section.left_margin = Inches(0.48)
        section.right_margin = Inches(0.48)

    for idx, t_data in enumerate(teacher_schedules):
        t_name = t_data['teacher']
        day_lessons = t_data['day_lessons']

        # 1. Meta Table
        meta_table = doc_out.add_table(rows=1, cols=2)
        meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        meta_cell_left = meta_table.cell(0, 0)
        meta_cell_right = meta_table.cell(0, 1)

        meta_cell_left.width = Inches(4.17)
        meta_cell_right.width = Inches(1.39)

        p_left = meta_cell_left.paragraphs[0]
        p_left.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r_left = p_left.add_run(META_TIMESTAMP)
        r_left.font.name = 'Arial'
        r_left.font.size = Pt(8.5)
        r_left.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

        p_right = meta_cell_right.paragraphs[0]
        p_right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_right = p_right.add_run(SCHOOL_NAME)
        r_right.font.name = 'Arial'
        r_right.font.size = Pt(8.5)
        r_right.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
        r_right.font.bold = True

        # 2. Title Paragraph
        p_title = doc_out.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pPr = p_title._p.get_or_add_pPr()
        pPr.append(parse_xml(r'<w:bidi %s/>' % nsdecls('w')))

        title_text = f"מערכת שעות מורה {t_name}".strip()
        r_title = p_title.add_run(title_text)
        r_title.font.name = 'Arial'
        r_title.font.size = Pt(14)
        r_title.font.bold = True
        r_title.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

        # 3. Schedule Table
        sched_table = doc_out.add_table(rows=0, cols=2)
        sched_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(sched_table)

        for day_num, day_name in DAYS_ORDER:
            lessons = day_lessons[day_num]
            if not lessons:
                continue

            lessons.sort(key=lambda x: x[0])

            # Day Header Row
            hdr_row = sched_table.add_row()
            c_left = hdr_row.cells[0]
            c_right = hdr_row.cells[1]

            c_left.width = Inches(4.17)
            c_right.width = Inches(1.39)

            set_cell_background(c_left, "28486B")
            set_cell_background(c_right, "28486B")

            p_lh = c_left.paragraphs[0]
            p_lh.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r_lh = p_lh.add_run("")
            r_lh.font.name = 'Arial'
            r_lh.font.size = Pt(9.0)

            p_rh = c_right.paragraphs[0]
            p_rh.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r_rh = p_rh.add_run(day_name)
            r_rh.font.name = 'Arial'
            r_rh.font.size = Pt(9.0)
            r_rh.font.bold = True
            r_rh.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

            set_cell_rtl(c_left)
            set_cell_rtl(c_right)

            # Lesson Rows
            for hour_num, content in lessons:
                l_row = sched_table.add_row()
                lc_left = l_row.cells[0]
                lc_right = l_row.cells[1]

                lc_left.width = Inches(4.17)
                lc_right.width = Inches(1.39)

                p_l = lc_left.paragraphs[0]
                p_l.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                r_l = p_l.add_run(content)
                r_l.font.name = 'Arial'
                r_l.font.size = Pt(9.0)
                r_l.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

                p_r = lc_right.paragraphs[0]
                p_r.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                lbl = HOUR_LABELS.get(hour_num, f"שעה {hour_num}")
                r_r = p_r.add_run(lbl)
                r_r.font.name = 'Arial'
                r_r.font.size = Pt(9.0)
                r_r.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

                set_cell_rtl(lc_left)
                set_cell_rtl(lc_right)

        # Page break between teachers
        if idx < len(teacher_schedules) - 1:
            p_break = doc_out.add_paragraph()
            p_break.add_run().add_break(WD_BREAK.PAGE)

    doc_out.save(TEACHER_DOCX)
    print(f"Successfully saved {TEACHER_DOCX}")


def map_class_pages(pdf_path):
    """Maps class_code -> (doc_object, page_index) from a class PDF."""
    doc = pymupdf.open(pdf_path)
    pages = {}
    for idx in range(len(doc)):
        p = doc[idx]
        lines = [l.strip() for l in p.get_text().split('\n') if l.strip()]
        title_line = [l for l in lines if 'מערכת שעות' in l][0]
        c_code = parse_class_title(title_line)
        pages[c_code] = (doc, idx)
    return pages


def create_classes_docx(ordered_classes, final_class_pages, class_from_teachers):
    print(f"Creating Classes DOCX: {CLASS_DOCX}")
    doc_out = Document()

    for section in doc_out.sections:
        section.top_margin = Inches(0.24)
        section.bottom_margin = Inches(0.48)
        section.left_margin = Inches(0.48)
        section.right_margin = Inches(0.48)

    for idx, c_code in enumerate(ordered_classes):
        doc_pdf, page_idx = final_class_pages[c_code]
        page = doc_pdf[page_idx]
        hr_teacher = CLASS_HOMEROOM.get(c_code, "")

        # 1. Meta Table
        meta_table = doc_out.add_table(rows=1, cols=2)
        meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        meta_cell_left = meta_table.cell(0, 0)
        meta_cell_right = meta_table.cell(0, 1)

        meta_cell_left.width = Inches(4.17)
        meta_cell_right.width = Inches(1.39)

        p_left = meta_cell_left.paragraphs[0]
        p_left.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r_left = p_left.add_run(META_TIMESTAMP)
        r_left.font.name = 'Arial'
        r_left.font.size = Pt(8.5)
        r_left.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

        p_right = meta_cell_right.paragraphs[0]
        p_right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_right = p_right.add_run(SCHOOL_NAME)
        r_right.font.name = 'Arial'
        r_right.font.size = Pt(8.5)
        r_right.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
        r_right.font.bold = True

        # 2. Title Paragraph
        p_title = doc_out.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pPr = p_title._p.get_or_add_pPr()
        pPr.append(parse_xml(r'<w:bidi %s/>' % nsdecls('w')))

        title_text = f"מערכת שעות כיתה {c_code} {hr_teacher}".strip()
        r_title = p_title.add_run(title_text)
        r_title.font.name = 'Arial'
        r_title.font.size = Pt(14)
        r_title.font.bold = True
        r_title.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

        # 3. Schedule Table
        sched_table = doc_out.add_table(rows=0, cols=2)
        sched_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(sched_table)

        # Parse days & hours for this class
        tabs = page.find_tables()
        if not tabs.tables:
            continue
        t = tabs.tables[0]

        day_lessons = {d_num: [] for d_num in range(1, 7)}

        for r_idx in range(2, len(t.rows)):
            row = t.rows[r_idx]
            h_cell = row.cells[6] if len(row.cells) > 6 else None
            h_text = page.get_text('text', clip=h_cell).strip() if h_cell else str(r_idx - 1)
            h_num = int(h_text) if h_text.isdigit() else (r_idx - 1)

            for col_idx, day_num in COL_TO_DAY.items():
                cell = row.cells[col_idx]
                if not cell:
                    continue
                txt = page.get_text('text', clip=cell).strip()
                if not txt:
                    continue

                # Get clean, canonical lessons for this slot from teachers
                lessons_from_t = class_from_teachers.get((c_code, day_num, h_num), [])
                if lessons_from_t:
                    # Format each lesson: "{Subject}, {Teacher}, {Type}"
                    lines = [f"{item['subject']}, {item['teacher']}, {item['type']}" for item in lessons_from_t]
                    content = "\n".join(lines)
                    day_lessons[day_num].append((h_num, content))
                else:
                    # Fallback to single lesson in cell (if NOT a workgroup)
                    raw_lines = [l.strip() for l in txt.split('\n') if l.strip()]
                    if len(raw_lines) >= 3:
                        l_type = raw_lines[0]
                        t_name = clean_teacher(raw_lines[1])
                        s_name = clean_subject(raw_lines[2])
                        if not is_workgroup(s_name) and not is_workgroup(l_type):
                            content = f"{s_name}, {t_name}, {l_type}"
                            day_lessons[day_num].append((h_num, content))

        for day_num, day_name in DAYS_ORDER:
            lessons = day_lessons[day_num]
            if not lessons:
                continue

            lessons.sort(key=lambda x: x[0])

            # Day Header Row
            hdr_row = sched_table.add_row()
            c_left = hdr_row.cells[0]
            c_right = hdr_row.cells[1]

            c_left.width = Inches(4.17)
            c_right.width = Inches(1.39)

            set_cell_background(c_left, "28486B")
            set_cell_background(c_right, "28486B")

            p_lh = c_left.paragraphs[0]
            p_lh.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r_lh = p_lh.add_run("")
            r_lh.font.name = 'Arial'
            r_lh.font.size = Pt(9.0)

            p_rh = c_right.paragraphs[0]
            p_rh.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r_rh = p_rh.add_run(day_name)
            r_rh.font.name = 'Arial'
            r_rh.font.size = Pt(9.0)
            r_rh.font.bold = True
            r_rh.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

            set_cell_rtl(c_left)
            set_cell_rtl(c_right)

            # Lesson Rows
            for hour_num, content in lessons:
                l_row = sched_table.add_row()
                lc_left = l_row.cells[0]
                lc_right = l_row.cells[1]

                lc_left.width = Inches(4.17)
                lc_right.width = Inches(1.39)

                p_l = lc_left.paragraphs[0]
                p_l.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                r_l = p_l.add_run(content)
                r_l.font.name = 'Arial'
                r_l.font.size = Pt(9.0)
                r_l.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

                p_r = lc_right.paragraphs[0]
                p_r.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                lbl = HOUR_LABELS.get(hour_num, f"שעה {hour_num}")
                r_r = p_r.add_run(lbl)
                r_r.font.name = 'Arial'
                r_r.font.size = Pt(9.0)
                r_r.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

                set_cell_rtl(lc_left)
                set_cell_rtl(lc_right)

        # Page break between classes
        if idx < len(ordered_classes) - 1:
            p_break = doc_out.add_paragraph()
            p_break.add_run().add_break(WD_BREAK.PAGE)

    doc_out.save(CLASS_DOCX)
    print(f"Successfully saved {CLASS_DOCX}")


def main():
    print("=== Starting Conversion for Yitzhak Sadeh A Bat Yam ===")
    
    # 1. Parse Base Teachers
    print(f"Loading Base Teachers PDF: {TEACHER_PDF}")
    base_scheds, base_order = extract_teachers_from_pdf(TEACHER_PDF)
    merged_teachers = dict(base_scheds)

    # 2. Overlay Teacher Changes from Batches
    for batch in BATCHES:
        t_pdf = batch["teacher_pdf"]
        if os.path.exists(t_pdf):
            print(f"Applying Teacher changes from {batch['name']}: {os.path.basename(t_pdf)}")
            b_scheds, _ = extract_teachers_from_pdf(t_pdf)
            for t_name, data in b_scheds.items():
                print(f"  -> Updated teacher: {t_name}")
                merged_teachers[t_name] = data

    # 3. Assemble Ordered Teacher Schedules and Class-From-Teachers
    final_teacher_schedules = []
    class_from_teachers = {}
    for t_name in base_order:
        data = merged_teachers[t_name]
        final_teacher_schedules.append({
            'teacher': t_name,
            'day_lessons': data['day_lessons']
        })
        for c_code, d_num, h_num, l_info in data['class_lessons']:
            key = (c_code, d_num, h_num)
            if key not in class_from_teachers:
                class_from_teachers[key] = []
            class_from_teachers[key].append(l_info)

    print(f"Total Teachers: {len(final_teacher_schedules)}")
    print(f"Total Class-Day-Hour Slots mapped: {len(class_from_teachers)}")

    # 4. Generate Teachers DOCX
    create_teachers_docx(final_teacher_schedules)

    # 5. Parse and Merge Classes
    print(f"\nLoading Base Classes PDF: {CLASS_PDF}")
    base_class_pages = map_class_pages(CLASS_PDF)

    doc_base = pymupdf.open(CLASS_PDF)
    ordered_classes = []
    for idx in range(len(doc_base)):
        p = doc_base[idx]
        lines = [l.strip() for l in p.get_text().split('\n') if l.strip()]
        title_line = [l for l in lines if 'מערכת שעות' in l][0]
        ordered_classes.append(parse_class_title(title_line))

    # Map batch pages
    batch_class_maps = []
    for batch in BATCHES:
        c_pdf = batch["class_pdf"]
        if os.path.exists(c_pdf):
            b_map = map_class_pages(c_pdf)
            batch_class_maps.append((batch["name"], b_map))

    final_class_pages = {}
    for c_code in ordered_classes:
        # Check batches in reverse order (newest first)
        found = False
        for b_name, b_map in reversed(batch_class_maps):
            if c_code in b_map:
                print(f"Using {c_code} from {b_name}")
                final_class_pages[c_code] = b_map[c_code]
                found = True
                break
        if not found:
            final_class_pages[c_code] = base_class_pages[c_code]

    # 6. Generate Classes DOCX
    create_classes_docx(ordered_classes, final_class_pages, class_from_teachers)
    print("=== Conversion Complete! ===")


if __name__ == '__main__':
    main()
