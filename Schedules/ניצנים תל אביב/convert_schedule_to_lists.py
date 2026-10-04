"""
Converter for "Nitzanim Tel Aviv" School Schedules
School ID: isfxl6jyrrhrmua8zkj9suyv

Converts:
  1) כיתות.pdf -> כיתות.docx
  2) מורים.pdf -> מורים.docx

Formatted in the exact Korczak style for seamless import into ShibutzPlus.
"""

import os
import sys
import re
from datetime import datetime
import pymupdf
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

if sys.stdout:
    sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
CLASS_PDF = os.path.join(BASE_DIR, "כיתות.pdf")
TEACHER_PDF = os.path.join(BASE_DIR, "מורים.pdf")
CLASS_DOCX = os.path.join(BASE_DIR, "כיתות.docx")
TEACHER_DOCX = os.path.join(BASE_DIR, "מורים.docx")

# School metadata
SCHOOL_ID = "isfxl6jyrrhrmua8zkj9suyv"
SCHOOL_NAME = "ניצנים"
META_TIMESTAMP = "04/10/2026 07:17:58"

HOUR_LABELS = {
    0: "שעה 0",
    1: "שעה 1",
    2: "שעה 2",
    3: "שעה 3",
    4: "שעה 4",
    5: "שעה 5",
    6: "שעה 6",
    7: "שעה 7",
    8: "שעה 8",
    9: "שעה 9",
    10: "שעה 10",
}

DAYS_ORDER = [
    (1, "יום ראשון"),
    (2, "יום שני"),
    (3, "יום שלישי"),
    (4, "יום רביעי"),
    (5, "יום חמישי"),
    (6, "יום שישי")
]

# Map PDF column index to Day number (1=Sunday, 6=Friday)
# Column 6: ראשון, 5: שני, 4: שלישי, 3: רביעי, 2: חמישי, 1: שישי
COL_TO_DAY = {
    6: 1,  # ראשון
    5: 2,  # שני
    4: 3,  # שלישי
    3: 4,  # רביעי
    2: 5,  # חמישי
    1: 6   # שישי
}

# Homeroom teachers mapping
CLASS_HOMEROOM = {
    'א1': 'מיכל כהן',
    'א2': 'גלית נאמן',
    'א3': 'אבי דרזיה',
    'ב1': 'נועה עייש שדה',
    'ב2': 'נועה עייש שדה',
    'ב3': 'ספיר אירינה',
    'ג1': 'לירון סכר',
    'ג2': 'אורין טוויזר',
    'ג3': 'שני וויס',
    'ד1': 'שני שטרית',
    'ד2': 'דנה טל',
    'ד3': 'אורי לוי',
    'ה1': 'שחף חרוש ארמי',
    'ה2': 'ריקי באגר',
    'ו1': 'אייל טל',
    'ו2': 'יובל פייסר',
    'ו3': 'מאיה לבל',
}

CLASS_DISPLAY_NAMES = {
    'א1': "כיתה א' 1",
    'א2': "כיתה א' 2",
    'א3': "כיתה א' 3",
    'ב1': "כיתה ב' 1",
    'ב2': "כיתה ב' 2",
    'ב3': "כיתה ב' 3",
    'ג1': "כיתה ג' 1",
    'ג2': "כיתה ג' 2",
    'ג3': "כיתה ג' 3",
    'ד1': "כיתה ד' 1",
    'ד2': "כיתה ד' 2",
    'ד3': "כיתה ד' 3",
    'ה1': "כיתה ה' 1",
    'ה2': "כיתה ה' 2",
    'ו1': "כיתה ו' 1",
    'ו2': "כיתה ו' 2",
    'ו3': "כיתה ו' 3",
}

# Map raw teacher names from PDF to canonical First Last
TEACHER_NAME_MAP = {
    'באגר ריקי': 'ריקי באגר',
    'זהבי מעין': 'מעין זהבי',
    'כהן עופרה': 'עופרה כהן',
    'לבל גלי': 'גלי לבל',
    'עייש שדה נועה': 'נועה עייש שדה',
    'פייסר יובל': 'יובל פייסר',
    'שטרית שני': 'שני שטרית',
    'שטרן עירית': 'עירית שטרן',
    'שן צור ליה': 'ליה שן צור',
    'אפי שדה': 'אפי שדה',
    'אפי': 'אפי שדה',
    'מדברים פלוס': 'אפי שדה',
    'רם גאן': 'רם גאן',
    'דור קרקובר': 'דור קרקובר',
    'דניאל סער': 'דניאל סער',
}

# WorkGroup / Individual activities
WORKGROUP_SUBJECTS = {
    'פרטני',
    'שהייה',
    'שהיית מליאה',
    'שעת תפקיד',
    'צוות ניהול',
    'צוות מגמות',
    'צוות אנגלית',
    'ריכוז בתים',
    'ריכוז בחירה',
    'ריכוז מגמות',
    'הדרכה מתמטיקה',
    'ליווי שפה',
    'מקהלה בוגרת',
    'מקהלה צעירה',
    'נבחרות',
    'תגבור',
    'העצמה',
    'צפייה',
    'פרטני שילוב',
    'רוחב ב2',
    'הוראה מותאמת',
    'בית א',
    'בית ב',
    'בית ג',
    'בית ד',
    'בית ה ו',
    'גאומטריה',
}

CLASS_CODE_REGEX = re.compile(r'([א-י][׳\']?[\s-]?[1-9][0-9]?|[א-י]["״][א-י][\s-]?[1-9][0-9]?)')

def normalize_class_code(raw: str) -> str:
    if not raw:
        return ""
    clean = raw.strip()
    m = CLASS_CODE_REGEX.search(clean)
    if m:
        code = m.group(1)
        code = re.sub(r'[\'\"״׳\u05F4\u05F3\u201C\u201D\u2018\u2019`\-]', '', code)
        code = re.sub(r'\s+', '', code)
        return code
    return ""

def is_workgroup(name: str) -> bool:
    if not name:
        return False
    name_clean = clean_subject(name)
    if name_clean in WORKGROUP_SUBJECTS:
        return True
    return any(k in name_clean for k in ['פרטני', 'שהייה', 'מליאה', 'תפקיד', 'צוות', 'ריכוז', 'הדרכ', 'מקהלה', 'נבחרות', 'תגבור', 'העצמה', 'צפייה', 'רוחב', 'שילוב', 'מותאמת', 'בית '])

def clean_subject(raw_subj: str) -> str:
    if not raw_subj:
        return ""
    s = raw_subj.strip()
    s = re.sub(r'[\u2018\u2019\u05F3\u00B4`]', "'", s)
    s = re.sub(r'[\u201C\u201D\u05F4]', '"', s)
    s = re.sub(r'\s+', ' ', s)
    # normalize quotes in תנך
    if s == 'תנך':
        s = 'תנ"ך'
    return s

def clean_teacher(name: str) -> str:
    if not name:
        return ""
    clean = re.sub(r'\s+', ' ', name).strip()
    return TEACHER_NAME_MAP.get(clean, clean)

def parse_teacher_title(title_line: str) -> str:
    clean = title_line.strip()
    clean = re.sub(r'מערכת שעות\s+(?:ל?מורה|מורה:?)\s*', '', clean).strip()
    return clean_teacher(clean)

def parse_class_title(title_line: str) -> str:
    # Handles e.g. "1 'מערכת שעות כיתה א" or "מערכת שעות כיתה א' 1"
    code = normalize_class_code(title_line)
    return code

def format_class_code_for_teacher(c_code: str) -> str:
    """Formats 'א1' -> 'א\' 1', 'ד3' -> 'ד\' 3'"""
    if len(c_code) >= 2:
        return f"{c_code[0]}' {c_code[1:]}"
    return c_code

def format_classes_string(raw_classes_str: str) -> str:
    """
    Parses class references from teacher cells.
    Handles 'ב' 2', '3'ד,2 'ד,1 'ד', etc.
    Returns: 'ד\' 1, ד\' 2, ד\' 3' or 'ב\' 2'
    """
    parts = [p.strip() for p in raw_classes_str.split(',') if p.strip()]
    formatted = []
    for p in parts:
        code = normalize_class_code(p)
        if code:
            formatted.append(format_class_code_for_teacher(code))
        else:
            formatted.append(p)
    return ", ".join(formatted)

def extract_class_codes_from_string(raw_classes_str: str) -> list[str]:
    parts = [p.strip() for p in raw_classes_str.split(',') if p.strip()]
    codes = []
    for p in parts:
        code = normalize_class_code(p)
        if code and code not in codes:
            codes.append(code)
    return codes

def set_cell_rtl(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(r'<w:tcMar %s><w:top w:w="80" w:type="dxa"/><w:bottom w:w="80" w:type="dxa"/><w:left w:w="120" w:type="dxa"/><w:right w:w="120" w:type="dxa"/></w:tcMar>' % nsdecls('w'))
    tcPr.append(tcMar)
    for p in cell.paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        pPr = p._p.get_or_add_pPr()
        pPr.append(parse_xml(r'<w:bidi %s/>' % nsdecls('w')))
        for r in p.runs:
            rPr = r._r.get_or_add_rPr()
            rPr.append(parse_xml(r'<w:rtl %s/>' % nsdecls('w')))

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="none" w:sz="4"/>'
        f'  <w:bottom w:val="none" w:sz="4"/>'
        f'  <w:left w:val="none" w:sz="4"/>'
        f'  <w:right w:val="none" w:sz="4"/>'
        f'  <w:insideH w:val="none" w:sz="4"/>'
        f'  <w:insideV w:val="none" w:sz="4"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def extract_data_from_teachers_pdf():
    """
    Parses TEACHER_PDF.
    Returns:
      teacher_schedules: list of dicts with teacher name and daily lessons
      class_from_teachers: (class_code, day, hour) -> list of {teacher, subject, type}
    """
    doc_pdf = pymupdf.open(TEACHER_PDF)
    teacher_schedules = []
    class_from_teachers = {}

    current_teacher = None
    teacher_entries = {}  # canonical_teacher -> {d_num: []}

    for pno in range(len(doc_pdf)):
        page = doc_pdf[pno]
        text = page.get_text()
        
        # Check if page defines a new teacher or continues previous
        t_match = None
        for line in text.split('\n'):
            if "מערכת שעות למורה" in line or "מערכת שעות מורה" in line:
                t_match = parse_teacher_title(line)
                break
                
        if t_match:
            current_teacher = t_match
            if current_teacher not in teacher_entries:
                teacher_entries[current_teacher] = {d_num: [] for d_num in range(1, 7)}
        elif not current_teacher:
            continue

        tabs = page.find_tables()
        if not tabs.tables:
            continue
        tab = tabs.tables[0]
        extracted = tab.extract()
        if len(extracted) < 3:
            continue

        for row in extracted[2:]:
            hour_val = row[7]
            if not hour_val or not hour_val.strip().isdigit():
                continue
            h_num = int(hour_val.strip())

            for col_idx, day_num in COL_TO_DAY.items():
                cell_text = row[col_idx]
                if not cell_text or not cell_text.strip():
                    continue

                # Reverse visual order Hebrew lines from PDF font
                c_lines = [l.strip()[::-1] for l in cell_text.split('\n') if l.strip()]

                if len(c_lines) == 1:
                    # Workgroup / Individual activity
                    wg_name = clean_subject(c_lines[0])
                    lesson_str = f"{wg_name}, קבוצה"
                    teacher_entries[current_teacher][day_num].append((h_num, lesson_str))

                elif len(c_lines) >= 3:
                    # Regular teaching lesson: Subject, Class(es), Type
                    subj = clean_subject(c_lines[0])
                    raw_cls = c_lines[1]
                    l_type = c_lines[2]

                    if is_workgroup(subj):
                        lesson_str = f"{subj}, קבוצה"
                    else:
                        cls_str = format_classes_string(raw_cls)
                        lesson_str = f"{subj}, {cls_str}, {l_type}"

                        # Map to classes
                        c_codes = extract_class_codes_from_string(raw_cls)
                        for c_code in c_codes:
                            key = (c_code, day_num, h_num)
                            if key not in class_from_teachers:
                                class_from_teachers[key] = []
                            # Avoid duplicate teacher entries for same slot
                            existing = [x['teacher'] for x in class_from_teachers[key]]
                            if current_teacher not in existing:
                                class_from_teachers[key].append({
                                    'teacher': current_teacher,
                                    'subject': subj,
                                    'type': l_type
                                })

                    teacher_entries[current_teacher][day_num].append((h_num, lesson_str))

    for t_name, day_lessons in teacher_entries.items():
        teacher_schedules.append({
            'teacher': t_name,
            'day_lessons': day_lessons
        })

    return teacher_schedules, class_from_teachers

def create_teachers_docx(teacher_schedules):
    print(f"Creating Teachers DOCX: {TEACHER_DOCX}")
    doc_out = Document()

    for section in doc_out.sections:
        section.top_margin = Inches(0.24)
        section.bottom_margin = Inches(0.48)
        section.left_margin = Inches(0.48)
        section.right_margin = Inches(0.48)

    for idx, t_data in enumerate(teacher_schedules):
        t_name = t_data['teacher']
        day_lessons = t_data['day_lessons']

        # 1. Meta Table
        meta_table = doc_out.add_table(rows=1, cols=2)
        meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        meta_cell_left = meta_table.cell(0, 0)
        meta_cell_right = meta_table.cell(0, 1)

        meta_cell_left.width = Inches(4.17)
        meta_cell_right.width = Inches(1.39)

        p_left = meta_cell_left.paragraphs[0]
        p_left.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r_left = p_left.add_run(META_TIMESTAMP)
        r_left.font.name = 'Arial'
        r_left.font.size = Pt(8.5)
        r_left.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

        p_right = meta_cell_right.paragraphs[0]
        p_right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_right = p_right.add_run(SCHOOL_NAME)
        r_right.font.name = 'Arial'
        r_right.font.size = Pt(8.5)
        r_right.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
        r_right.font.bold = True

        # 2. Title Paragraph
        p_title = doc_out.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pPr = p_title._p.get_or_add_pPr()
        pPr.append(parse_xml(r'<w:bidi %s/>' % nsdecls('w')))

        title_text = f"מערכת שעות מורה {t_name}".strip()
        r_title = p_title.add_run(title_text)
        r_title.font.name = 'Arial'
        r_title.font.size = Pt(14)
        r_title.font.bold = True
        r_title.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

        # 3. Schedule Table
        sched_table = doc_out.add_table(rows=0, cols=2)
        sched_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(sched_table)

        for day_num, day_name in DAYS_ORDER:
            lessons = day_lessons[day_num]
            if not lessons:
                continue

            lessons.sort(key=lambda x: x[0])

            # Day Header Row
            hdr_row = sched_table.add_row()
            c_left = hdr_row.cells[0]
            c_right = hdr_row.cells[1]

            c_left.width = Inches(4.17)
            c_right.width = Inches(1.39)

            set_cell_background(c_left, "28486B")
            set_cell_background(c_right, "28486B")

            p_lh = c_left.paragraphs[0]
            p_lh.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r_lh = p_lh.add_run("")
            r_lh.font.name = 'Arial'
            r_lh.font.size = Pt(9.0)

            p_rh = c_right.paragraphs[0]
            p_rh.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r_rh = p_rh.add_run(day_name)
            r_rh.font.name = 'Arial'
            r_rh.font.size = Pt(9.0)
            r_rh.font.bold = True
            r_rh.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

            set_cell_rtl(c_left)
            set_cell_rtl(c_right)

            # Lesson Rows
            for hour_num, content in lessons:
                l_row = sched_table.add_row()
                lc_left = l_row.cells[0]
                lc_right = l_row.cells[1]

                lc_left.width = Inches(4.17)
                lc_right.width = Inches(1.39)

                p_l = lc_left.paragraphs[0]
                p_l.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                r_l = p_l.add_run(content)
                r_l.font.name = 'Arial'
                r_l.font.size = Pt(9.0)
                r_l.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

                p_r = lc_right.paragraphs[0]
                p_r.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                lbl = HOUR_LABELS.get(hour_num, f"שעה {hour_num}")
                r_r = p_r.add_run(lbl)
                r_r.font.name = 'Arial'
                r_r.font.size = Pt(9.0)
                r_r.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

                set_cell_rtl(lc_left)
                set_cell_rtl(lc_right)

        # Page break between teachers
        if idx < len(teacher_schedules) - 1:
            p_break = doc_out.add_paragraph()
            p_break.add_run().add_break(WD_BREAK.PAGE)

    doc_out.save(TEACHER_DOCX)
    print(f"Successfully saved {TEACHER_DOCX}")

def create_classes_docx(class_from_teachers):
    print(f"Creating Classes DOCX: {CLASS_DOCX}")
    doc_c = pymupdf.open(CLASS_PDF)
    doc_out = Document()

    for section in doc_out.sections:
        section.top_margin = Inches(0.24)
        section.bottom_margin = Inches(0.48)
        section.left_margin = Inches(0.48)
        section.right_margin = Inches(0.48)

    for idx in range(len(doc_c)):
        page = doc_c[idx]
        tabs = page.find_tables()
        if not tabs.tables:
            continue
        tab = tabs.tables[0]
        extracted = tab.extract()
        if len(extracted) < 3:
            continue

        raw_title = extracted[0][0][::-1] if extracted[0][0] else ""
        c_code = parse_class_title(raw_title)
        display_name = CLASS_DISPLAY_NAMES.get(c_code, f"כיתה {c_code}")
        hr_teacher = CLASS_HOMEROOM.get(c_code, "")

        # 1. Meta Table
        meta_table = doc_out.add_table(rows=1, cols=2)
        meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        meta_cell_left = meta_table.cell(0, 0)
        meta_cell_right = meta_table.cell(0, 1)

        meta_cell_left.width = Inches(4.17)
        meta_cell_right.width = Inches(1.39)

        p_left = meta_cell_left.paragraphs[0]
        p_left.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r_left = p_left.add_run(META_TIMESTAMP)
        r_left.font.name = 'Arial'
        r_left.font.size = Pt(8.5)
        r_left.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

        p_right = meta_cell_right.paragraphs[0]
        p_right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_right = p_right.add_run(SCHOOL_NAME)
        r_right.font.name = 'Arial'
        r_right.font.size = Pt(8.5)
        r_right.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
        r_right.font.bold = True

        # 2. Title Paragraph
        p_title = doc_out.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pPr = p_title._p.get_or_add_pPr()
        pPr.append(parse_xml(r'<w:bidi %s/>' % nsdecls('w')))

        title_text = f"מערכת שעות {display_name}".strip()
        r_title = p_title.add_run(title_text)
        r_title.font.name = 'Arial'
        r_title.font.size = Pt(14)
        r_title.font.bold = True
        r_title.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

        # 3. Schedule Table
        sched_table = doc_out.add_table(rows=0, cols=2)
        sched_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(sched_table)

        day_lessons = {d_num: [] for d_num in range(1, 7)}

        for row in extracted[2:]:
            hour_val = row[7]
            if not hour_val or not hour_val.strip().isdigit():
                continue
            h_num = int(hour_val.strip())

            for col_idx, day_num in COL_TO_DAY.items():
                cell_text = row[col_idx]
                if not cell_text or not cell_text.strip():
                    continue

                # Check if we have authoritative lessons from teachers
                lessons_from_t = class_from_teachers.get((c_code, day_num, h_num), [])
                if lessons_from_t:
                    lines = [f"{item['subject']}, {item['teacher']}, {item['type']}" for item in lessons_from_t]
                    content = "\n".join(lines)
                    day_lessons[day_num].append((h_num, content))
                else:
                    # Fallback to class cell
                    raw_lines = [l.strip()[::-1] for l in cell_text.split('\n') if l.strip()]
                    if len(raw_lines) >= 3:
                        subj = clean_subject(raw_lines[0])
                        teach = clean_teacher(raw_lines[1])
                        l_type = raw_lines[2]
                        content = f"{subj}, {teach}, {l_type}"
                        day_lessons[day_num].append((h_num, content))

        for day_num, day_name in DAYS_ORDER:
            lessons = day_lessons[day_num]
            if not lessons:
                continue

            lessons.sort(key=lambda x: x[0])

            # Day Header Row
            hdr_row = sched_table.add_row()
            c_left = hdr_row.cells[0]
            c_right = hdr_row.cells[1]

            c_left.width = Inches(4.17)
            c_right.width = Inches(1.39)

            set_cell_background(c_left, "28486B")
            set_cell_background(c_right, "28486B")

            p_lh = c_left.paragraphs[0]
            p_lh.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r_lh = p_lh.add_run("")
            r_lh.font.name = 'Arial'
            r_lh.font.size = Pt(9.0)

            p_rh = c_right.paragraphs[0]
            p_rh.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r_rh = p_rh.add_run(day_name)
            r_rh.font.name = 'Arial'
            r_rh.font.size = Pt(9.0)
            r_rh.font.bold = True
            r_rh.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

            set_cell_rtl(c_left)
            set_cell_rtl(c_right)

            # Lesson Rows
            for hour_num, content in lessons:
                l_row = sched_table.add_row()
                lc_left = l_row.cells[0]
                lc_right = l_row.cells[1]

                lc_left.width = Inches(4.17)
                lc_right.width = Inches(1.39)

                p_l = lc_left.paragraphs[0]
                p_l.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                r_l = p_l.add_run(content)
                r_l.font.name = 'Arial'
                r_l.font.size = Pt(9.0)
                r_l.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

                p_r = lc_right.paragraphs[0]
                p_r.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                lbl = HOUR_LABELS.get(hour_num, f"שעה {hour_num}")
                r_r = p_r.add_run(lbl)
                r_r.font.name = 'Arial'
                r_r.font.size = Pt(9.0)
                r_r.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

                set_cell_rtl(lc_left)
                set_cell_rtl(lc_right)

        # Page break between classes
        if idx < len(doc_c) - 1:
            p_break = doc_out.add_paragraph()
            p_break.add_run().add_break(WD_BREAK.PAGE)

    doc_out.save(CLASS_DOCX)
    print(f"Successfully saved {CLASS_DOCX}")

def main():
    print("=== Starting Conversion for Nitzanim Tel Aviv ===")
    teacher_schedules, class_from_teachers = extract_data_from_teachers_pdf()
    print(f"Extracted {len(teacher_schedules)} teacher schedules.")
    print(f"Extracted {len(class_from_teachers)} class-day-hour slots.")

    create_teachers_docx(teacher_schedules)
    create_classes_docx(class_from_teachers)
    print("=== Conversion Complete! ===")

if __name__ == '__main__':
    main()
